import os
import sys
import glob
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

# Thử import thư viện Kéo - Thả
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    HAS_DND = False


# Từ điển đa ngôn ngữ (Multilingual Dictionary - 16 Languages)
TRANSLATIONS = {
    "English": {
        "title": "RPYtoEXCEL",
        "language_label": "Language:",
        "tab_export": " 1. Export RPY -> Excel ",
        "tab_merge": " 2. Merge Excel -> RPY ",
        "export_frame": " Extract Translation Files ",
        "export_desc": "Select or DROP .rpy files / folders into the list below:",
        "btn_add_files": "Add RPY Files...",
        "btn_add_folder": "Add Folder...",
        "btn_clear_list": "Clear List",
        "btn_start_export": " START EXPORT TO EXCEL ",
        "merge_frame": " Merge Translations into RPY Files ",
        "lbl_excel_step": "1. Select or DROP translated Excel file here:",
        "btn_browse_excel": "Browse Excel...",
        "lbl_rpy_step": "2. Select or DROP target .rpy file(s) / folder:",
        "btn_start_merge": " START MERGING TRANSLATIONS ",
        "excel_notes_sheet": "Please provide translations in the 'Translated' column. Once finished, save the file and use the 'Merge Excel -> RPY' tab to apply translations to the game.",
        "msg_warn_select_rpy": "Please select at least one .rpy file!",
        "msg_success_export": "File successfully exported to:\n{}",
        "msg_err_export": "An error occurred while exporting the Excel file:\n{}",
        "msg_warn_excel": "Please select a valid Excel file!",
        "msg_warn_rpy_folder": "Please select at least one .rpy file or directory!",
        "msg_success_merge": "Successfully merged translations into {} .rpy file(s)!",
        "msg_err_merge": "An error occurred while merging:\n{}",
        "dlg_select_rpy": "Select .rpy Files",
        "dlg_save_excel": "Save Output Excel File",
        "dlg_select_excel": "Select Translated Excel File",
        "dlg_select_folder": "Select Folder Containing .rpy Files",
        "warning": "Warning",
        "success": "Success",
        "error": "Error",
        "col_file": "File",
        "col_line": "Line",
        "col_location": "Original Location",
        "col_label": "Label",
        "col_tag": "Tag Char",
        "col_original": "Original",
        "col_translated": "Translated",
        "col_notes": "Notes",
        "sheet_notes_name": "Notes",
        "sheet_addition_name": "Addition",
        "progress_title": "Progress",
        "sum_title": "Sum"
    },
    "Tiếng Việt": {
        "title": "RPYtoEXCEL",
        "language_label": "Ngôn ngữ / Language:",
        "tab_export": " 1. Trích Xuất RPY -> Excel ",
        "tab_merge": " 2. Ghép Excel -> RPY ",
        "export_frame": " Trích xuất file dịch ",
        "export_desc": "Chọn hoặc KÉO THẢ các file/thư mục .rpy vào danh sách:",
        "btn_add_files": "Thêm File RPY...",
        "btn_add_folder": "Thêm Thư Mục...",
        "btn_clear_list": "Xóa Danh Sách",
        "btn_start_export": " BẮT ĐẦU XUẤT EXCEL ",
        "merge_frame": " Ghép bản dịch vào File RPY ",
        "lbl_excel_step": "1. Chọn hoặc KÉO THẢ file Excel đã dịch vào ô bên dưới:",
        "btn_browse_excel": "Chọn Excel...",
        "lbl_rpy_step": "2. Chọn hoặc KÉO THẢ các file/Thư mục .rpy cần gộp:",
        "btn_start_merge": " BẮT ĐẦU GỘP BẢN DỊCH ",
        "excel_notes_sheet": "Vui lòng dịch ở cột 'Bản dịch'. Sau khi hoàn tất, lưu file và sử dụng tab 'Ghép Excel -> RPY' để hợp nhất bản dịch vào game.",
        "msg_warn_select_rpy": "Vui lòng chọn ít nhất 1 file .rpy!",
        "msg_success_export": "Đã xuất file thành công tại:\n{}",
        "msg_err_export": "Có lỗi xảy ra khi xuất file Excel:\n{}",
        "msg_warn_excel": "Vui lòng chọn file Excel bản dịch hợp lệ!",
        "msg_warn_rpy_folder": "Vui lòng chọn ít nhất 1 file hoặc thư mục chứa .rpy!",
        "msg_success_merge": "Đã hợp nhất thành công bản dịch vào {} file .rpy!",
        "msg_err_merge": "Có lỗi xảy ra khi merge:\n{}",
        "dlg_select_rpy": "Chọn các file .rpy",
        "dlg_save_excel": "Lưu file Excel kết quả",
        "dlg_select_excel": "Chọn file Excel bản dịch",
        "dlg_select_folder": "Chọn thư mục chứa các file .rpy cần cập nhật",
        "warning": "Cảnh báo",
        "success": "Thành công",
        "error": "Lỗi",
        "col_file": "Tên File",
        "col_line": "Dòng",
        "col_location": "Vị trí gốc",
        "col_label": "Nhãn (Label)",
        "col_tag": "Thẻ nhân vật",
        "col_original": "Bản gốc",
        "col_translated": "Bản dịch",
        "col_notes": "Ghi chú",
        "sheet_notes_name": "Lưu ý!",
        "sheet_addition_name": "Bổ sung",
        "progress_title": "Tiến độ",
        "sum_title": "Tổng cộng"
    },
    "Español": {
        "title": "RPYtoEXCEL",
        "language_label": "Idioma / Language:",
        "tab_export": " 1. Exportar RPY -> Excel ",
        "tab_merge": " 2. Combinar Excel -> RPY ",
        "export_frame": " Extraer archivos de traducción ",
        "export_desc": "Seleccione o ARRASTRE archivos/carpetas .rpy aquí:",
        "btn_add_files": "Añadir archivos RPY...",
        "btn_add_folder": "Añadir carpeta...",
        "btn_clear_list": "Limpiar lista",
        "btn_start_export": " INICIAR EXPORTACIÓN ",
        "merge_frame": " Combinar traducciones en archivos RPY ",
        "lbl_excel_step": "1. Seleccione o ARRASTRE el archivo Excel traducido:",
        "btn_browse_excel": "Buscar Excel...",
        "lbl_rpy_step": "2. Seleccione o ARRASTRE archivos .rpy o carpeta:",
        "btn_start_merge": " INICIAR COMBINACIÓN ",
        "excel_notes_sheet": "Traduzca en la columna 'Traducción'. Guarde el archivo y use la pestaña 'Combinar Excel -> RPY'.",
        "msg_warn_select_rpy": "¡Seleccione al menos un archivo .rpy!",
        "msg_success_export": "Archivo exportado con éxito en:\n{}",
        "msg_err_export": "Error al exportar el archivo Excel:\n{}",
        "msg_warn_excel": "¡Seleccione un archivo Excel válido!",
        "msg_warn_rpy_folder": "¡Seleccione al menos un archivo .rpy o carpeta!",
        "msg_success_merge": "¡Traducciones combinadas con éxito en {} archivo(s)!",
        "msg_err_merge": "Error al combinar:\n{}",
        "dlg_select_rpy": "Seleccionar archivos .rpy",
        "dlg_save_excel": "Guardar archivo Excel",
        "dlg_select_excel": "Seleccionar archivo Excel",
        "dlg_select_folder": "Seleccionar carpeta .rpy",
        "warning": "Advertencia",
        "success": "Éxito",
        "error": "Error",
        "col_file": "Archivo",
        "col_line": "Línea",
        "col_location": "Ubicación original",
        "col_label": "Etiqueta",
        "col_tag": "Etiqueta Personaje",
        "col_original": "Original",
        "col_translated": "Traducción",
        "col_notes": "Notas",
        "sheet_notes_name": "Notas",
        "sheet_addition_name": "Adicional",
        "progress_title": "Progreso",
        "sum_title": "Total"
    },
    "中文": {
        "title": "RPYtoEXCEL",
        "language_label": "语言 / Language:",
        "tab_export": " 1. 导出 RPY -> Excel ",
        "tab_merge": " 2. 合并 Excel -> RPY ",
        "export_frame": " 提取翻译文件 ",
        "export_desc": "选择或拖放 .rpy 文件/文件夹到下方列表：",
        "btn_add_files": "添加 RPY 文件...",
        "btn_add_folder": "添加文件夹...",
        "btn_clear_list": "清空列表",
        "btn_start_export": " 开始导出 Excel ",
        "merge_frame": " 将翻译合并到 RPY 文件 ",
        "lbl_excel_step": "1. 选择或拖放已翻译的 Excel 文件：",
        "btn_browse_excel": "浏览 Excel...",
        "lbl_rpy_step": "2. 选择或拖放目标 .rpy 文件或文件夹：",
        "btn_start_merge": " 开始合并翻译 ",
        "excel_notes_sheet": "请在 '译文' 列中提供翻译。完成后保存并使用 '合并 Excel -> RPY' 选项卡。",
        "msg_warn_select_rpy": "请至少选择一个 .rpy 文件！",
        "msg_success_export": "成功导出至：\n{}",
        "msg_err_export": "导出 Excel 文件时出错：\n{}",
        "msg_warn_excel": "请选择有效的 Excel 文件！",
        "msg_warn_rpy_folder": "请选择至少一个 .rpy 文件或目录！",
        "msg_success_merge": "成功合并翻译至 {} 个 .rpy 文件！",
        "msg_err_merge": "合并时出错：\n{}",
        "dlg_select_rpy": "选择 .rpy 文件",
        "dlg_save_excel": "保存输出 Excel 文件",
        "dlg_select_excel": "选择翻译好的 Excel 文件",
        "dlg_select_folder": "选择包含 .rpy 文件的文件夹",
        "warning": "警告",
        "success": "成功",
        "error": "错误",
        "col_file": "文件名",
        "col_line": "行号",
        "col_location": "原始位置",
        "col_label": "标签",
        "col_tag": "角色标签",
        "col_original": "原文",
        "col_translated": "译文",
        "col_notes": "备注",
        "sheet_notes_name": "备注",
        "sheet_addition_name": "补充",
        "progress_title": "进度",
        "sum_title": "总计"
    },
    "Deutsch": {
        "title": "RPYtoEXCEL",
        "language_label": "Sprache / Language:",
        "tab_export": " 1. Exportieren RPY -> Excel ",
        "tab_merge": " 2. Zusammenführen Excel -> RPY ",
        "export_frame": " Übersetzungsdateien extrahieren ",
        "export_desc": "Wählen Sie .rpy-Dateien/Ordner aus oder ziehen Sie sie hierher:",
        "btn_add_files": "RPY-Dateien hinzufügen...",
        "btn_add_folder": "Ordner hinzufügen...",
        "btn_clear_list": "Liste leeren",
        "btn_start_export": " EXPORT STARTEN ",
        "merge_frame": " Übersetzungen in RPY-Dateien zusammenführen ",
        "lbl_excel_step": "1. Übersetzte Excel-Datei auswählen oder hierher ziehen:",
        "btn_browse_excel": "Excel durchsuchen...",
        "lbl_rpy_step": "2. Ziel-.rpy-Dateien oder Ordner auswählen/ziehen:",
        "btn_start_merge": " ZUSAMMENFÜHRUNG STARTEN ",
        "excel_notes_sheet": "Tragen Sie Übersetzungen in die Spalte 'Übersetzung' ein.",
        "msg_warn_select_rpy": "Bitte wählen Sie mindestens eine .rpy-Datei aus!",
        "msg_success_export": "Erfolgreich exportiert nach:\n{}",
        "msg_err_export": "Fehler beim Exportieren der Excel-Datei:\n{}",
        "msg_warn_excel": "Bitte wählen Sie eine gültige Excel-Datei aus!",
        "msg_warn_rpy_folder": "Bitte wählen Sie mindestens eine Datei oder einen Ordner aus!",
        "msg_success_merge": "Erfolgreich in {} .rpy-Datei(en) zusammengeführt!",
        "msg_err_merge": "Fehler beim Zusammenführen:\n{}",
        "dlg_select_rpy": ".rpy-Dateien auswählen",
        "dlg_save_excel": "Excel-Datei speichern",
        "dlg_select_excel": "Excel-Datei auswählen",
        "dlg_select_folder": "Ordner auswählen",
        "warning": "Warnung",
        "success": "Erfolg",
        "error": "Fehler",
        "col_file": "Datei",
        "col_line": "Zeile",
        "col_location": "Originaler Pfad",
        "col_label": "Label",
        "col_tag": "Charakter-Tag",
        "col_original": "Original",
        "col_translated": "Übersetzung",
        "col_notes": "Notizen",
        "sheet_notes_name": "Notizen",
        "sheet_addition_name": "Zusatz",
        "progress_title": "Fortschritt",
        "sum_title": "Gesamt"
    },
    "Français": {
        "title": "RPYtoEXCEL",
        "language_label": "Langue / Language:",
        "tab_export": " 1. Exporter RPY -> Excel ",
        "tab_merge": " 2. Fusionner Excel -> RPY ",
        "export_frame": " Extraire les fichiers de traduction ",
        "export_desc": "Sélectionnez ou glissez-déposez les fichiers/dossiers .rpy :",
        "btn_add_files": "Ajouter fichiers RPY...",
        "btn_add_folder": "Ajouter dossier...",
        "btn_clear_list": "Effacer la liste",
        "btn_start_export": " LANCER L'EXPORTATION ",
        "merge_frame": " Fusionner dans les fichiers RPY ",
        "lbl_excel_step": "1. Sélectionner ou glisser le fichier Excel traduit :",
        "btn_browse_excel": "Parcourir Excel...",
        "lbl_rpy_step": "2. Sélectionner ou glisser les fichiers .rpy ou le dossier :",
        "btn_start_merge": " LANCER LA FUSION ",
        "excel_notes_sheet": "Traduisez dans la colonne 'Traduction'.",
        "msg_warn_select_rpy": "Veuillez sélectionner au moins un fichier .rpy !",
        "msg_success_export": "Exporté avec succès vers :\n{}",
        "msg_err_export": "Erreur lors de l'exportation :\n{}",
        "msg_warn_excel": "Veuillez sélectionner un fichier Excel valide !",
        "msg_warn_rpy_folder": "Veuillez sélectionner au moins un fichier ou un dossier !",
        "msg_success_merge": "Fusion réussie dans {} fichier(s) .rpy !",
        "msg_err_merge": "Erreur lors de la fusion :\n{}",
        "dlg_select_rpy": "Sélectionner fichiers .rpy",
        "dlg_save_excel": "Enregistrer le fichier Excel",
        "dlg_select_excel": "Sélectionner le fichier Excel",
        "dlg_select_folder": "Sélectionner le dossier",
        "warning": "Avertissement",
        "success": "Succès",
        "error": "Erreur",
        "col_file": "Fichier",
        "col_line": "Ligne",
        "col_location": "Emplacement d'origine",
        "col_label": "Étiquette",
        "col_tag": "Tag Personnage",
        "col_original": "Original",
        "col_translated": "Traduction",
        "col_notes": "Notes",
        "sheet_notes_name": "Notes",
        "sheet_addition_name": "Addition",
        "progress_title": "Progression",
        "sum_title": "Total"
    },
    "Bahasa Indonesia": {
        "title": "RPYtoEXCEL",
        "language_label": "Bahasa / Language:",
        "tab_export": " 1. Ekspor RPY -> Excel ",
        "tab_merge": " 2. Gabung Excel -> RPY ",
        "export_frame": " Ekstrak Berkas Terjemahan ",
        "export_desc": "Pilih atau TARIK berkas/folder .rpy ke daftar di bawah:",
        "btn_add_files": "Tambah Berkas RPY...",
        "btn_add_folder": "Tambah Folder...",
        "btn_clear_list": "Bersihkan Daftar",
        "btn_start_export": " MULAI EKSPOR ",
        "merge_frame": " Gabung Terjemahan ke Berkas RPY ",
        "lbl_excel_step": "1. Pilih atau TARIK berkas Excel di sini:",
        "btn_browse_excel": "Cari Excel...",
        "lbl_rpy_step": "2. Pilih atau TARIK berkas .rpy / folder tujuan:",
        "btn_start_merge": " MULAI PENGGABUNGAN ",
        "excel_notes_sheet": "Isi terjemahan di kolom 'Terjemahan'.",
        "msg_warn_select_rpy": "Pilih setidaknya satu berkas .rpy!",
        "msg_success_export": "Berhasil diekspor ke:\n{}",
        "msg_err_export": "Galat saat mengekspor berkas Excel:\n{}",
        "msg_warn_excel": "Pilih berkas Excel yang valid!",
        "msg_warn_rpy_folder": "Pilih setidaknya satu berkas atau folder!",
        "msg_success_merge": "Berhasil menggabungkan ke {} berkas .rpy!",
        "msg_err_merge": "Galat saat menggabungkan:\n{}",
        "dlg_select_rpy": "Pilih Berkas .rpy",
        "dlg_save_excel": "Simpan Berkas Excel",
        "dlg_select_excel": "Pilih Berkas Excel",
        "dlg_select_folder": "Pilih Folder",
        "warning": "Peringatan",
        "success": "Berhasil",
        "error": "Galat",
        "col_file": "Nama Berkas",
        "col_line": "Baris",
        "col_location": "Lokasi Asli",
        "col_label": "Label",
        "col_tag": "Tag Karakter",
        "col_original": "Teks Asli",
        "col_translated": "Terjemahan",
        "col_notes": "Catatan",
        "sheet_notes_name": "Catatan",
        "sheet_addition_name": "Tambahan",
        "progress_title": "Kemajuan",
        "sum_title": "Total"
    },
    "العربية": {
        "title": "RPYtoEXCEL",
        "language_label": "اللغة / Language:",
        "tab_export": " 1. تصدير RPY -> Excel ",
        "tab_merge": " 2. دمج Excel -> RPY ",
        "export_frame": " استخراج ملفات الترجمة ",
        "export_desc": "حدد أو اسحب ملفات/مجلدات .rpy أدناه:",
        "btn_add_files": "إضافة ملفات RPY...",
        "btn_add_folder": "إضافة مجلد...",
        "btn_clear_list": "مسح القائمة",
        "btn_start_export": " بدء التصدير ",
        "merge_frame": " دمج الترجمات في ملفات RPY ",
        "lbl_excel_step": "1. حدد أو اسحب ملف Excel المترجم هنا:",
        "btn_browse_excel": "تصفح Excel...",
        "lbl_rpy_step": "2. حدد أو اسحب ملفات .rpy أو المجلد الهدف:",
        "btn_start_merge": " بدء الدمج ",
        "excel_notes_sheet": "يرجى توفير الترجمات في عمود الترجمة.",
        "msg_warn_select_rpy": "يرجى تحديد ملف .rpy واحد على الأقل!",
        "msg_success_export": "تم التصدير بنجاح إلى:\n{}",
        "msg_err_export": "حدث خطأ أثناء تصدير Excel:\n{}",
        "msg_warn_excel": "يرجى تحديد ملف Excel صالحة!",
        "msg_warn_rpy_folder": "يرجى تحديد ملف واحد على الأقل أو مجلد!",
        "msg_success_merge": "تم الدمج بنجاح في {} ملف(ملفات) .rpy!",
        "msg_err_merge": "حدث خطأ أثناء الدمج:\n{}",
        "dlg_select_rpy": "حدد ملفات .rpy",
        "dlg_save_excel": "حفظ ملف Excel",
        "dlg_select_excel": "حدد ملف Excel",
        "dlg_select_folder": "حدد المجلد",
        "warning": "تحذير",
        "success": "نجاح",
        "error": "خطأ",
        "col_file": "اسم الملف",
        "col_line": "السطر",
        "col_location": "الموقع الأصلي",
        "col_label": "العلامة",
        "col_tag": "علامة الشخصية",
        "col_original": "النص الأصلي",
        "col_translated": "الترجمة",
        "col_notes": "ملاحظات",
        "sheet_notes_name": "ملاحظات",
        "sheet_addition_name": "إضافة",
        "progress_title": "التقدم",
        "sum_title": "المجموع"
    },
    "Türkçe": {
        "title": "RPYtoEXCEL",
        "language_label": "Dil / Language:",
        "tab_export": " 1. Dışa Aktar RPY -> Excel ",
        "tab_merge": " 2. Birleştir Excel -> RPY ",
        "export_frame": " Çeviri Dosyalarını Ayıkla ",
        "export_desc": ".rpy dosyalarını/klasörlerini seçin veya buraya sürükleyin:",
        "btn_add_files": "RPY Dosyası Ekle...",
        "btn_add_folder": "Klasör Ekle...",
        "btn_clear_list": "Listeyi Temizle",
        "btn_start_export": " DIŞA AKTARMAYI BAŞLAT ",
        "merge_frame": " Çevirileri RPY Dosyalarına Birleştir ",
        "lbl_excel_step": "1. Çevrilmiş Excel dosyasını seçin veya sürükleyin:",
        "btn_browse_excel": "Excel Gözat...",
        "lbl_rpy_step": "2. Hedef .rpy dosyalarını/klasörü seçin veya sürükleyin:",
        "btn_start_merge": " BİRLEŞTİRMEYİ BAŞLAT ",
        "excel_notes_sheet": "Lütfen çevirileri 'Çeviri' sütununa yazın.",
        "msg_warn_select_rpy": "Lütfen en az bir .rpy dosyası seçin!",
        "msg_success_export": "Başarıyla şuraya aktarıldı:\n{}",
        "msg_err_export": "Excel aktarılırken bir hata oluştu:\n{}",
        "msg_warn_excel": "Lütfen geçerli bir Excel dosyası seçin!",
        "msg_warn_rpy_folder": "Lütfen en az bir dosyası veya klasör seçin!",
        "msg_success_merge": "Çeviriler {} .rpy dosyasına başarıyla birleştirildi!",
        "msg_err_merge": "Birleştirilirken hata oluştu:\n{}",
        "dlg_select_rpy": ".rpy Dosyalarını Seç",
        "dlg_save_excel": "Excel Dosyasını Kaydet",
        "dlg_select_excel": "Excel Dosyasını Seç",
        "dlg_select_folder": "Klasör Seç",
        "warning": "Uyarı",
        "success": "Başarılı",
        "error": "Hata",
        "col_file": "Dosya Adı",
        "col_line": "Satır",
        "col_location": "Orijinal Konum",
        "col_label": "Etiket",
        "col_tag": "Karakter Etiketi",
        "col_original": "Orijinal Metin",
        "col_translated": "Çeviri",
        "col_notes": "Notlar",
        "sheet_notes_name": "Notlar",
        "sheet_addition_name": "Ek",
        "progress_title": "İlerleme",
        "sum_title": "Toplam"
    },
    "Português": {
        "title": "RPYtoEXCEL",
        "language_label": "Idioma / Language:",
        "tab_export": " 1. Exportar RPY -> Excel ",
        "tab_merge": " 2. Mesclar Excel -> RPY ",
        "export_frame": " Extrair arquivos de tradução ",
        "export_desc": "Selecione ou ARRASTE arquivos/pastas .rpy para a lista:",
        "btn_add_files": "Adicionar arquivos RPY...",
        "btn_add_folder": "Adicionar pasta...",
        "btn_clear_list": "Limpar lista",
        "btn_start_export": " INICIAR EXPORTAÇÃO ",
        "merge_frame": " Mesclar traduções nos arquivos RPY ",
        "lbl_excel_step": "1. Selecione ou ARRASTE o arquivo Excel traduzido:",
        "btn_browse_excel": "Procurar Excel...",
        "lbl_rpy_step": "2. Selecione ou ARRASTE arquivos .rpy ou pasta:",
        "btn_start_merge": " INICIAR MESCLAGEM ",
        "excel_notes_sheet": "Traduza na coluna 'Tradução'.",
        "msg_warn_select_rpy": "Selecione pelo menos um arquivo .rpy!",
        "msg_success_export": "Exportado com sucesso para:\n{}",
        "msg_err_export": "Erro ao exportar arquivo Excel:\n{}",
        "msg_warn_excel": "Selecione um arquivo Excel válido!",
        "msg_warn_rpy_folder": "Selecione pelo menos um arquivo ou pasta!",
        "msg_success_merge": "Traduções mescladas com sucesso em {} arquivo(s) .rpy!",
        "msg_err_merge": "Erro ao mesclar:\n{}",
        "dlg_select_rpy": "Selecionar arquivos .rpy",
        "dlg_save_excel": "Salvar arquivo Excel",
        "dlg_select_excel": "Selecionar arquivo Excel",
        "dlg_select_folder": "Selecionar pasta",
        "warning": "Aviso",
        "success": "Sucesso",
        "error": "Erro",
        "col_file": "Nome do Arquivo",
        "col_line": "Linha",
        "col_location": "Localização Original",
        "col_label": "Rótulo",
        "col_tag": "Tag Personagem",
        "col_original": "Texto Original",
        "col_translated": "Tradução",
        "col_notes": "Notas",
        "sheet_notes_name": "Notas",
        "sheet_addition_name": "Adição",
        "progress_title": "Progresso",
        "sum_title": "Total"
    },
    "Polski": {
        "title": "RPYtoEXCEL",
        "language_label": "Język / Language:",
        "tab_export": " 1. Eksportuj RPY -> Excel ",
        "tab_merge": " 2. Scal Excel -> RPY ",
        "export_frame": " Wyodrębnij pliki tłumaczeń ",
        "export_desc": "Wybierz lub PRZECIĄGNIJ pliki/foldery .rpy do listy:",
        "btn_add_files": "Dodaj pliki RPY...",
        "btn_add_folder": "Dodaj folder...",
        "btn_clear_list": "Wyczyść listę",
        "btn_start_export": " ROZPOCZNIJ EKSPORT ",
        "merge_frame": " Scal tłumaczenia z plikami RPY ",
        "lbl_excel_step": "1. Wybierz lub PRZECIĄGNIJ przetłumaczony plik Excel:",
        "btn_browse_excel": "Przeglądaj Excel...",
        "lbl_rpy_step": "2. Wybierz lub PRZECIĄGNIJ docelowe pliki .rpy lub folder:",
        "btn_start_merge": " ROZPOCZNIJ SCALANIE ",
        "excel_notes_sheet": "Wpisz tłumaczenia w kolumnie 'Tłumaczenie'.",
        "msg_warn_select_rpy": "Wybierz co najmniej jeden plik .rpy!",
        "msg_success_export": "Pomyślnie wyeksportowano do:\n{}",
        "msg_err_export": "Błąd podczas eksportowania pliku Excel:\n{}",
        "msg_warn_excel": "Wybierz prawidłowy plik Excel!",
        "msg_warn_rpy_folder": "Wybierz co najmniej jeden plik lub folder!",
        "msg_success_merge": "Pomyślnie scalono tłumaczenia w {} plikach .rpy!",
        "msg_err_merge": "Błąd podczas scalania:\n{}",
        "dlg_select_rpy": "Wybierz pliki .rpy",
        "dlg_save_excel": "Zapisz plik Excel",
        "dlg_select_excel": "Wybierz plik Excel",
        "dlg_select_folder": "Wybierz folder",
        "warning": "Ostrzeżenie",
        "success": "Sukces",
        "error": "Błąd",
        "col_file": "Nazwa Pliku",
        "col_line": "Linia",
        "col_location": "Oryginalna Lokalizacja",
        "col_label": "Etykieta",
        "col_tag": "Tag Postaci",
        "col_original": "Oryginał",
        "col_translated": "Tłumaczenie",
        "col_notes": "Notatki",
        "sheet_notes_name": "Notatki",
        "sheet_addition_name": "Dodatek",
        "progress_title": "Postęp",
        "sum_title": "Suma"
    },
    "Русский": {
        "title": "RPYtoEXCEL",
        "language_label": "Язык / Language:",
        "tab_export": " 1. Экспорт RPY -> Excel ",
        "tab_merge": " 2. Слияние Excel -> RPY ",
        "export_frame": " Извлечь файлы перевода ",
        "export_desc": "Выберите или ПЕРЕТАЩИТЕ файлы/папки .rpy:",
        "btn_add_files": "Добавить файлы RPY...",
        "btn_add_folder": "Добавить папку...",
        "btn_clear_list": "Очистить список",
        "btn_start_export": " НАЧАТЬ ЭКСПОРТ В EXCEL ",
        "merge_frame": " Объединить переводы с файлами RPY ",
        "lbl_excel_step": "1. Выберите или ПЕРЕТАЩИТЕ файл Excel:",
        "btn_browse_excel": "Обзор Excel...",
        "lbl_rpy_step": "2. Выберите или ПЕРЕТАЩИТЕ файлы .rpy или папку:",
        "btn_start_merge": " НАЧАТЬ СЛИЯНИЕ ПЕРЕВОДОВ ",
        "excel_notes_sheet": "Укажите перевод в столбце 'Перевод'.",
        "msg_warn_select_rpy": "Выберите хотя бы один файл .rpy!",
        "msg_success_export": "Успешно экспортировано в:\n{}",
        "msg_err_export": "Ошибка при экспорте файла Excel:\n{}",
        "msg_warn_excel": "Выберите корректный файл Excel!",
        "msg_warn_rpy_folder": "Выберите хотя бы один файл или папку!",
        "msg_success_merge": "Успешно объединены переводы в {} файл(ах) .rpy!",
        "msg_err_merge": "Ошибка при объединении:\n{}",
        "dlg_select_rpy": "Выберите файлы .rpy",
        "dlg_save_excel": "Сохранить файл Excel",
        "dlg_select_excel": "Выберите файл Excel",
        "dlg_select_folder": "Выберите папку с файлами .rpy",
        "warning": "Предупреждение",
        "success": "Успех",
        "error": "Ошибка",
        "col_file": "Имя файла",
        "col_line": "Строка",
        "col_location": "Исходное место",
        "col_label": "Метка",
        "col_tag": "Тег персонажа",
        "col_original": "Оригинал",
        "col_translated": "Перевод",
        "col_notes": "Заметки",
        "sheet_notes_name": "Заметки",
        "sheet_addition_name": "Дополнение",
        "progress_title": "Прогресс",
        "sum_title": "Итого"
    },
    "Українська": {
        "title": "RPYtoEXCEL",
        "language_label": "Мова / Language:",
        "tab_export": " 1. Експорт RPY -> Excel ",
        "tab_merge": " 2. Злиття Excel -> RPY ",
        "export_frame": " Витягти файли перекладу ",
        "export_desc": "Виберіть або ПЕРЕТЯГНІТЬ файли/папки .rpy нижче:",
        "btn_add_files": "Додати файли RPY...",
        "btn_add_folder": "Додати папку...",
        "btn_clear_list": "Очистити список",
        "btn_start_export": " ПОЧАТИ ЕКСПОРТ ",
        "merge_frame": " Об'єднати переклади з файлами RPY ",
        "lbl_excel_step": "1. Виберіть або ПЕРЕТЯГНІТЬ перекладений файл Excel:",
        "btn_browse_excel": "Огляд Excel...",
        "lbl_rpy_step": "2. Виберіть або ПЕРЕТЯГНІТЬ файли .rpy або папку:",
        "btn_start_merge": " ПОЧАТИ ЗЛИТТЯ ",
        "excel_notes_sheet": "Вкажіть переклад у стовпчику 'Переклад'.",
        "msg_warn_select_rpy": "Виберіть принаймні один файл .rpy!",
        "msg_success_export": "Успішно експортовано в:\n{}",
        "msg_err_export": "Помилка під час експорту Excel:\n{}",
        "msg_warn_excel": "Виберіть коректний файл Excel!",
        "msg_warn_rpy_folder": "Виберіть принаймні один файл або папку!",
        "msg_success_merge": "Успішно об'єднано переклади в {} файл(ах) .rpy!",
        "msg_err_merge": "Помилка під час злиття:\n{}",
        "dlg_select_rpy": "Виберіть файли .rpy",
        "dlg_save_excel": "Зберегти файл Excel",
        "dlg_select_excel": "Виберіть файл Excel",
        "dlg_select_folder": "Виберіть папку з файлами .rpy",
        "warning": "Попередження",
        "success": "Успіх",
        "error": "Помилка",
        "col_file": "Назва файлу",
        "col_line": "Рядок",
        "col_location": "Оригінальне місце",
        "col_label": "Мітка",
        "col_tag": "Тег персонажа",
        "col_original": "Оригінал",
        "col_translated": "Переклад",
        "col_notes": "Примітки",
        "sheet_notes_name": "Примітки",
        "sheet_addition_name": "Доповнення",
        "progress_title": "Прогрес",
        "sum_title": "Всього"
    },
    "日本語": {
        "title": "RPYtoEXCEL",
        "language_label": "言語 / Language:",
        "tab_export": " 1. 抽出 RPY -> Excel ",
        "tab_merge": " 2. 統合 Excel -> RPY ",
        "export_frame": " 翻訳ファイルの抽出 ",
        "export_desc": ".rpy ファイル/フォルダを選択またはドラッグ＆ドロップしてください:",
        "btn_add_files": "RPY ファイルを追加...",
        "btn_add_folder": "フォルダを追加...",
        "btn_clear_list": "リストをクリア",
        "btn_start_export": " EXCEL 抽出を開始 ",
        "merge_frame": " 翻訳を RPY ファイルに統合 ",
        "lbl_excel_step": "1. 翻訳済み Excel ファイルを選択またはドロップ:",
        "btn_browse_excel": "Excel を参照...",
        "lbl_rpy_step": "2. 対象の .rpy ファイル/フォルダを選択またはドロップ:",
        "btn_start_merge": " 翻訳の統合を開始 ",
        "excel_notes_sheet": "'翻訳' 列に翻訳を入力してください。",
        "msg_warn_select_rpy": "少なくとも1つの .rpy ファイルを選択してください！",
        "msg_success_export": "正常に抽出されました:\n{}",
        "msg_err_export": "Excel の抽出中にエラーが発生しました:\n{}",
        "msg_warn_excel": "有効な Excel ファイルを選択してください！",
        "msg_warn_rpy_folder": "少なくとも1つのファイルまたはフォルダを選択してください！",
        "msg_success_merge": "{} 個の .rpy ファイルに翻訳を正常に統合しました！",
        "msg_err_merge": "統合中にエラーが発生しました:\n{}",
        "dlg_select_rpy": ".rpy ファイルを選択",
        "dlg_save_excel": "Excel ファイルを保存",
        "dlg_select_excel": "Excel ファイルを選択",
        "dlg_select_folder": "フォルダを選択",
        "warning": "警告",
        "success": "成功",
        "error": "エラー",
        "col_file": "ファイル名",
        "col_line": "行番号",
        "col_location": "元の場所",
        "col_label": "ラベル",
        "col_tag": "キャラタグ",
        "col_original": "原文",
        "col_translated": "翻訳",
        "col_notes": "メモ",
        "sheet_notes_name": "メモ",
        "sheet_addition_name": "追加",
        "progress_title": "進捗率",
        "sum_title": "合計"
    },
    "한국어": {
        "title": "RPYtoEXCEL",
        "language_label": "언어 / Language:",
        "tab_export": " 1. 추출 RPY -> Excel ",
        "tab_merge": " 2. 병합 Excel -> RPY ",
        "export_frame": " 번역 파일 추출 ",
        "export_desc": ".rpy 파일이나 폴더를 선택하거나 끌어다 놓으세요:",
        "btn_add_files": "RPY 파일 추가...",
        "btn_add_folder": "폴더 추가...",
        "btn_clear_list": "목록 비우기",
        "btn_start_export": " EXCEL 추출 시작 ",
        "merge_frame": " RPY 파일에 번역 병합 ",
        "lbl_excel_step": "1. 번역된 Excel 파일을 선택하거나 끌어다 놓으세요:",
        "btn_browse_excel": "Excel 찾아보기...",
        "lbl_rpy_step": "2. 대상 .rpy 파일 또는 폴더를 선택/끌어다 놓으세요:",
        "btn_start_merge": " 번역 병합 시작 ",
        "excel_notes_sheet": "'번역' 열에 번역문을 입력하세요.",
        "msg_warn_select_rpy": "최소 하나의 .rpy 파일을 선택하세요!",
        "msg_success_export": "성공적으로 내보냈습니다:\n{}",
        "msg_err_export": "Excel 파일 내보내기 중 오류 발생:\n{}",
        "msg_warn_excel": "올바른 Excel 파일을 선택하세요!",
        "msg_warn_rpy_folder": "최소 하나의 파일 또는 폴더를 선택하세요!",
        "msg_success_merge": "{}개의 .rpy 파일에 번역을 성공적으로 병합했습니다!",
        "msg_err_merge": "병합 중 오류 발생:\n{}",
        "dlg_select_rpy": ".rpy 파일 선택",
        "dlg_save_excel": "Excel 파일 저장",
        "dlg_select_excel": "Excel 파일 선택",
        "dlg_select_folder": "폴더 선택",
        "warning": "경고",
        "success": "성공",
        "error": "오류",
        "col_file": "파일명",
        "col_line": "줄 번호",
        "col_location": "원래 위치",
        "col_label": "라벨",
        "col_tag": "캐릭터 태그",
        "col_original": "원문",
        "col_translated": "번역",
        "col_notes": "메모",
        "sheet_notes_name": "메모",
        "sheet_addition_name": "추가",
        "progress_title": "진행률",
        "sum_title": "합계"
    },
    "Italiano": {
        "title": "RPYtoEXCEL",
        "language_label": "Lingua / Language:",
        "tab_export": " 1. Esporta RPY -> Excel ",
        "tab_merge": " 2. Unisci Excel -> RPY ",
        "export_frame": " Estrai file di traduzione ",
        "export_desc": "Seleziona o TRASCINA i file/cartelle .rpy qui:",
        "btn_add_files": "Aggiungi file RPY...",
        "btn_add_folder": "Aggiungi cartella...",
        "btn_clear_list": "Svuota lista",
        "btn_start_export": " AVVIA ESPORTAZIONE ",
        "merge_frame": " Unisci traduzioni nei file RPY ",
        "lbl_excel_step": "1. Seleziona o TRASCINA il file Excel tradotto:",
        "btn_browse_excel": "Sfoglia Excel...",
        "lbl_rpy_step": "2. Seleziona o TRASCINA file .rpy o cartella:",
        "btn_start_merge": " AVVIA UNIONE TRADUZIONI ",
        "excel_notes_sheet": "Inserisci le traduzioni nella colonna 'Traduzione'.",
        "msg_warn_select_rpy": "Seleziona almeno un file .rpy!",
        "msg_success_export": "Esportato con successo in:\n{}",
        "msg_err_export": "Errore durante l'esportazione Excel:\n{}",
        "msg_warn_excel": "Seleziona un file Excel valido!",
        "msg_warn_rpy_folder": "Seleziona almeno un file o cartella!",
        "msg_success_merge": "Traduzioni unite con successo in {} file .rpy!",
        "msg_err_merge": "Errore durante l'unione:\n{}",
        "dlg_select_rpy": "Seleziona file .rpy",
        "dlg_save_excel": "Salva file Excel",
        "dlg_select_excel": "Seleziona file Excel",
        "dlg_select_folder": "Seleziona cartella",
        "warning": "Avviso",
        "success": "Successo",
        "error": "Errore",
        "col_file": "Nome File",
        "col_line": "Riga",
        "col_location": "Posizione Originale",
        "col_label": "Etichetta",
        "col_tag": "Tag Personaggio",
        "col_original": "Originale",
        "col_translated": "Traduzione",
        "col_notes": "Note",
        "sheet_notes_name": "Note",
        "sheet_addition_name": "Aggiunta",
        "progress_title": "Progresso",
        "sum_title": "Totale"
    }
}

class RPYtoEXCELApp:
    def __init__(self, root):
        self.root = root
        self.current_lang = "Tiếng Việt"
        
        self.root.title("RPYtoEXCEL")
        self.root.geometry("680x580")
        self.root.resizable(True, True)

        style = ttk.Style()
        style.theme_use('clam')

        lang_frame = ttk.Frame(self.root)
        lang_frame.pack(fill='x', padx=15, pady=(10, 0))

        self.lbl_lang_select = ttk.Label(lang_frame, text="Language:", font=('Segoe UI', 9, 'bold'))
        self.lbl_lang_select.pack(side='left', padx=(0, 5))

        self.lang_var = tk.StringVar(value=self.current_lang)
        self.lang_menu = ttk.OptionMenu(
            lang_frame, 
            self.lang_var, 
            self.current_lang, 
            *list(TRANSLATIONS.keys()), 
            command=self.change_language
        )
        self.lang_menu.pack(side='left')

        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill='both', expand=True, padx=10, pady=10)

        self.tab_export = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_export, text="")
        self.setup_export_tab()

        self.tab_merge = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_merge, text="")
        self.setup_merge_tab()

        self.update_ui_text()

    def change_language(self, selected_lang):
        self.current_lang = selected_lang
        self.update_ui_text()

    def t(self, key):
        return TRANSLATIONS.get(self.current_lang, TRANSLATIONS["English"]).get(key, key)

    def update_ui_text(self):
        self.root.title(self.t("title"))
        self.lbl_lang_select.config(text=self.t("language_label"))
        
        self.notebook.tab(self.tab_export, text=self.t("tab_export"))
        self.notebook.tab(self.tab_merge, text=self.t("tab_merge"))

        self.export_lf.config(text=self.t("export_frame"))
        self.lbl_export_desc.config(text=self.t("export_desc"))
        self.btn_add.config(text=self.t("btn_add_files"))
        self.btn_add_folder.config(text=self.t("btn_add_folder"))
        self.btn_clear.config(text=self.t("btn_clear_list"))
        self.btn_export.config(text=self.t("btn_start_export"))

        self.merge_lf.config(text=self.t("merge_frame"))
        self.lbl_excel_step.config(text=self.t("lbl_excel_step"))
        self.btn_browse_excel.config(text=self.t("btn_browse_excel"))
        self.lbl_rpy_step.config(text=self.t("lbl_rpy_step"))
        self.btn_merge_add_files.config(text=self.t("btn_add_files"))
        self.btn_merge_add_folder.config(text=self.t("btn_add_folder"))
        self.btn_merge_clear.config(text=self.t("btn_clear_list"))
        self.btn_merge.config(text=self.t("btn_start_merge"))

    def setup_export_tab(self):
        self.export_lf = ttk.LabelFrame(self.tab_export, text="")
        self.export_lf.pack(fill='both', expand=True, padx=15, pady=15)

        self.lbl_export_desc = ttk.Label(self.export_lf, text="", font=('Segoe UI', 10))
        self.lbl_export_desc.pack(anchor='w', padx=10, pady=5)

        list_frame = ttk.Frame(self.export_lf)
        list_frame.pack(fill='both', expand=True, padx=10, pady=5)

        self.export_listbox = tk.Listbox(list_frame, selectmode=tk.EXTENDED, height=8)
        self.export_listbox.pack(side='left', fill='both', expand=True)

        scrollbar = ttk.Scrollbar(list_frame, orient='vertical', command=self.export_listbox.yview)
        scrollbar.pack(side='right', fill='y')
        self.export_listbox.config(yscrollcommand=scrollbar.set)

        if HAS_DND:
            self.export_listbox.drop_target_register(DND_FILES)
            self.export_listbox.dnd_bind('<<Drop>>', self.drop_export_files)

        btn_box = ttk.Frame(self.export_lf)
        btn_box.pack(fill='x', padx=10, pady=5)

        self.btn_add = ttk.Button(btn_box, text="", command=self.add_rpy_files)
        self.btn_add.pack(side='left', padx=5)

        self.btn_add_folder = ttk.Button(btn_box, text="", command=self.add_rpy_folder)
        self.btn_add_folder.pack(side='left', padx=5)

        self.btn_clear = ttk.Button(btn_box, text="", command=lambda: self.export_listbox.delete(0, tk.END))
        self.btn_clear.pack(side='left', padx=5)

        out_frame = ttk.Frame(self.export_lf)
        out_frame.pack(fill='x', padx=10, pady=10)

        self.btn_export = ttk.Button(out_frame, text="", command=self.process_export)
        self.btn_export.pack(fill='x', ipady=5)

    def drop_export_files(self, event):
        files = self.root.tk.splitlist(event.data)
        for path in files:
            path = path.strip('{}')
            if os.path.isfile(path) and path.endswith('.rpy'):
                if path not in self.export_listbox.get(0, tk.END):
                    self.export_listbox.insert(tk.END, path)
            elif os.path.isdir(path):
                rpy_files = glob.glob(os.path.join(path, "**", "*.rpy"), recursive=True)
                for f in rpy_files:
                    if f not in self.export_listbox.get(0, tk.END):
                        self.export_listbox.insert(tk.END, f)

    def add_rpy_files(self):
        files = filedialog.askopenfilenames(
            title=self.t("dlg_select_rpy"),
            filetypes=[("Ren'Py Script", "*.rpy"), ("All Files", "*.*")]
        )
        for f in files:
            if f not in self.export_listbox.get(0, tk.END):
                self.export_listbox.insert(tk.END, f)

    def add_rpy_folder(self):
        d = filedialog.askdirectory(title=self.t("dlg_select_folder"))
        if d:
            rpy_files = glob.glob(os.path.join(d, "**", "*.rpy"), recursive=True)
            for f in rpy_files:
                if f not in self.export_listbox.get(0, tk.END):
                    self.export_listbox.insert(tk.END, f)

    def process_export(self):
        rpy_files = list(self.export_listbox.get(0, tk.END))
        if not rpy_files:
            messagebox.showwarning(self.t("warning"), self.t("msg_warn_select_rpy"))
            return

        save_path = filedialog.asksaveasfilename(
            title=self.t("dlg_save_excel"),
            defaultextension=".xlsx",
            initialfile="Translation.xlsx",
            filetypes=[("Excel Files", "*.xlsx")]
        )
        if not save_path:
            return

        columns = [
            self.t("col_file"),
            self.t("col_line"),
            self.t("col_location"),
            self.t("col_label"),
            self.t("col_tag"),
            self.t("col_original"),
            self.t("col_translated"),
            self.t("col_notes")
        ]

        notes_sheet_title = self.t("sheet_notes_name")
        addition_sheet_title = self.t("sheet_addition_name")

        try:
            with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
                df_notes = pd.DataFrame([self.t("excel_notes_sheet")])
                df_notes.to_excel(writer, sheet_name=notes_sheet_title, index=False)

                table_mapping = []

                for table_idx, file_path in enumerate(rpy_files, start=1):
                    file_name = os.path.basename(file_path)
                    if not os.path.exists(file_path):
                        continue
                        
                    with open(file_path, "r", encoding="utf-8") as f:
                        lines = f.readlines()
                        
                    current_game_location = ""
                    current_label = ""
                    extracted_data = []
                        
                    for idx, line in enumerate(lines, start=1):
                        line_strip = line.strip()
                        
                        if line_strip.startswith("# game/") and ":" in line_strip:
                            current_game_location = line_strip.replace("#", "").strip()
                            
                        if line_strip.startswith("translate ") and line_strip.endswith(":"):
                            parts = line_strip.split()
                            if len(parts) >= 3:
                                current_label = parts[2].replace(":", "").strip()
                        
                        if line_strip.startswith("#") and '"' in line_strip:
                            start_pos = line.find('"')
                            end_pos = line.rfind('"')
                            
                            if start_pos != end_pos:
                                english_text = line[start_pos + 1 : end_pos]
                                if english_text.startswith("game/") or english_text.endswith(".rpy"):
                                    continue
                                
                                character_name = ""
                                translated_text = ""
                                
                                for offset in range(1, 4):
                                    next_idx = (idx - 1) + offset
                                    if next_idx >= len(lines):
                                        break
                                        
                                    next_line_strip = lines[next_idx].strip()
                                    if next_line_strip.startswith("#") or next_line_strip.startswith("translate"):
                                        break
                                    
                                    if next_line_strip.endswith('""'):
                                        possible_name = next_line_strip[:-2].strip()
                                        if possible_name:
                                            character_name = possible_name
                                    
                                    elif '"' in next_line_strip and not next_line_strip.startswith("#"):
                                        t_start = next_line_strip.find('"')
                                        t_end = next_line_strip.rfind('"')
                                        if t_start != t_end:
                                            translated_text = next_line_strip[t_start + 1 : t_end]
                                            if not character_name:
                                                tag = next_line_strip[:t_start].strip()
                                                if '"' not in tag:
                                                    character_name = tag
                                            break
                                
                                if not character_name:
                                    tag = line_strip[1:start_pos].strip()
                                    if '"' in tag or tag.startswith("game/"):
                                        character_name = ""
                                    else:
                                        character_name = tag

                                extracted_data.append([file_name, idx, current_game_location, current_label, character_name, english_text, translated_text, ""])

                        elif line_strip.startswith("old") and '"' in line_strip:
                            start_pos = line.find('"')
                            end_pos = line.rfind('"')
                            
                            if start_pos != end_pos:
                                english_text = line[start_pos + 1 : end_pos]
                                translated_text = ""
                                
                                for offset in range(1, 3):
                                    next_idx = (idx - 1) + offset
                                    if next_idx >= len(lines):
                                        break
                                    
                                    next_line_strip = lines[next_idx].strip()
                                    if next_line_strip.startswith("new") and '"' in next_line_strip:
                                        t_start = next_line_strip.find('"')
                                        t_end = next_line_strip.rfind('"')
                                        if t_start != t_end:
                                            translated_text = next_line_strip[t_start + 1 : t_end]
                                            break

                                extracted_data.append([file_name, idx, current_game_location, "strings", "", english_text, translated_text, ""])

                    if extracted_data:
                        df = pd.DataFrame(extracted_data, columns=columns)
                        sheet_name = file_name.replace(".rpy", "")[:30]
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        
                        table_name = f"Bảng_{table_idx}"
                        table_mapping.append((sheet_name, table_name))

                # --- Tạo Sheet Addition (Bổ sung) chuẩn định dạng Excel Table ---
                df_addition = pd.DataFrame(columns=columns)
                df_addition.to_excel(writer, sheet_name=addition_sheet_title, index=False)
                addition_table_name = f"Bảng_Addition"
                table_mapping.append((addition_sheet_title, addition_table_name))

                workbook = writer.book

                green_fill = PatternFill(start_color="356854", end_color="356854", fill_type="solid")
                gray_fill = PatternFill(start_color="F6F8F9", end_color="F6F8F9", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")

                thin_border = Border(
                    left=Side(style='thin', color='D3D3D3'),
                    right=Side(style='thin', color='D3D3D3'),
                    top=Side(style='thin', color='D3D3D3'),
                    bottom=Side(style='thin', color='D3D3D3')
                )

                for sheet_name, table_name in table_mapping:
                    ws = workbook[sheet_name]
                    max_row = max(ws.max_row, 2)
                    tab = Table(displayName=table_name, ref=f"A1:H{max_row}")
                    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
                    tab.tableStyleInfo = style
                    ws.add_table(tab)

                    for col in range(1, 9):
                        cell = ws.cell(row=1, column=col)
                        cell.fill = green_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

                    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=2):
                        for cell in row:
                            cell.border = thin_border
                            if row_idx % 2 == 0:
                                cell.fill = gray_fill
                                
                            if cell.column_letter in ['F', 'G', 'H']:
                                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
                            else:
                                cell.alignment = Alignment(vertical="center", horizontal="center")

                    fixed_widths = {'F': 55, 'G': 55, 'H': 25}
                    for col in ws.columns:
                        col_letter = col[0].column_letter
                        if col_letter in fixed_widths:
                            ws.column_dimensions[col_letter].width = fixed_widths[col_letter]
                        else:
                            max_len = 0
                            for cell in col:
                                val_str = str(cell.value or '')
                                val_len = len(val_str) + 6 if cell.row == 1 else len(val_str)
                                if val_len > max_len:
                                    max_len = val_len
                            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                # --- Lập báo cáo Tiến độ ở sheet Notes (LOẠI BỎ SHEET ADDITION KHỎI TIẾN ĐỘ) ---
                ws_notes = workbook[notes_sheet_title]
                cell_c3 = ws_notes.cell(row=3, column=3, value=self.t("col_file"))
                cell_d3 = ws_notes.cell(row=3, column=4, value=self.t("progress_title"))
                
                for cell in [cell_c3, cell_d3]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = green_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                col_trans_name = self.t("col_translated")
                col_orig_name = self.t("col_original")
                max_filename_len = len(self.t("col_file"))

                start_row = 4
                progress_table_mapping = [item for item in table_mapping if item[0] != addition_sheet_title]

                for idx, (sheet_name, table_name) in enumerate(progress_table_mapping, start=start_row):
                    cell_file = ws_notes.cell(row=idx, column=3, value=sheet_name)
                    cell_file.alignment = Alignment(horizontal="center", vertical="center")
                    
                    if len(sheet_name) > max_filename_len:
                        max_filename_len = len(sheet_name)
                    
                    cell_prog = ws_notes.cell(row=idx, column=4)
                    cell_prog.value = f"=COUNTA({table_name}[{col_trans_name}])/COUNTA({table_name}[{col_orig_name}])"
                    cell_prog.number_format = '0.00%'
                    cell_prog.alignment = Alignment(horizontal="center", vertical="center")

                sum_row = start_row + len(progress_table_mapping)
                cell_sum_lbl = ws_notes.cell(row=sum_row, column=3, value=self.t("sum_title"))
                cell_sum_lbl.font = Font(bold=True)
                cell_sum_lbl.alignment = Alignment(horizontal="center", vertical="center")
                
                if progress_table_mapping:
                    trans_parts = [f"COUNTA({tname}[{col_trans_name}])" for _, tname in progress_table_mapping]
                    orig_parts = [f"COUNTA({tname}[{col_orig_name}])" for _, tname in progress_table_mapping]
                    sum_formula = f"=({' + '.join(trans_parts)})/({' + '.join(orig_parts)})"
                    
                    cell_sum = ws_notes.cell(row=sum_row, column=4)
                    cell_sum.value = sum_formula
                    cell_sum.font = Font(bold=True)
                    cell_sum.number_format = '0.00%'
                    cell_sum.alignment = Alignment(horizontal="center", vertical="center")

                ws_notes.column_dimensions["C"].width = max(max_filename_len + 8, 25)
                ws_notes.column_dimensions["D"].width = max(len(self.t("progress_title")) + 8, 18)

            messagebox.showinfo(self.t("success"), self.t("msg_success_export").format(save_path))

        except Exception as e:
            messagebox.showerror(self.t("error"), self.t("msg_err_export").format(str(e)))

    def setup_merge_tab(self):
        self.merge_lf = ttk.LabelFrame(self.tab_merge, text="")
        self.merge_lf.pack(fill='both', expand=True, padx=15, pady=15)

        self.lbl_excel_step = ttk.Label(self.merge_lf, text="", font=('Segoe UI', 10, 'bold'))
        self.lbl_excel_step.pack(anchor='w', padx=10, pady=(5, 2))

        excel_box = ttk.Frame(self.merge_lf)
        excel_box.pack(fill='x', padx=10, pady=2)

        self.entry_excel = ttk.Entry(excel_box)
        self.entry_excel.pack(side='left', fill='x', expand=True, padx=(0, 5))

        self.btn_browse_excel = ttk.Button(excel_box, text="", command=self.browse_excel_file)
        self.btn_browse_excel.pack(side='right')

        if HAS_DND:
            self.entry_excel.drop_target_register(DND_FILES)
            self.entry_excel.dnd_bind('<<Drop>>', self.drop_excel_file)

        self.lbl_rpy_step = ttk.Label(self.merge_lf, text="", font=('Segoe UI', 10, 'bold'))
        self.lbl_rpy_step.pack(anchor='w', padx=10, pady=(10, 2))

        list_frame = ttk.Frame(self.merge_lf)
        list_frame.pack(fill='both', expand=True, padx=10, pady=5)

        self.merge_listbox = tk.Listbox(list_frame, selectmode=tk.EXTENDED, height=6)
        self.merge_listbox.pack(side='left', fill='both', expand=True)

        scrollbar = ttk.Scrollbar(list_frame, orient='vertical', command=self.merge_listbox.yview)
        scrollbar.pack(side='right', fill='y')
        self.merge_listbox.config(yscrollcommand=scrollbar.set)

        if HAS_DND:
            self.merge_listbox.drop_target_register(DND_FILES)
            self.merge_listbox.dnd_bind('<<Drop>>', self.drop_merge_files)

        btn_box = ttk.Frame(self.merge_lf)
        btn_box.pack(fill='x', padx=10, pady=5)

        self.btn_merge_add_files = ttk.Button(btn_box, text="", command=self.add_merge_rpy_files)
        self.btn_merge_add_files.pack(side='left', padx=5)

        self.btn_merge_add_folder = ttk.Button(btn_box, text="", command=self.add_merge_rpy_folder)
        self.btn_merge_add_folder.pack(side='left', padx=5)

        self.btn_merge_clear = ttk.Button(btn_box, text="", command=lambda: self.merge_listbox.delete(0, tk.END))
        self.btn_merge_clear.pack(side='left', padx=5)

        self.btn_merge = ttk.Button(self.merge_lf, text="", command=self.process_merge)
        self.btn_merge.pack(fill='x', padx=10, pady=10, ipady=5)

    def drop_excel_file(self, event):
        files = self.root.tk.splitlist(event.data)
        if files:
            file_path = files[0].strip('{}')
            if file_path.endswith(('.xlsx', '.xls')):
                self.entry_excel.delete(0, tk.END)
                self.entry_excel.insert(0, file_path)

    def drop_merge_files(self, event):
        files = self.root.tk.splitlist(event.data)
        for path in files:
            path = path.strip('{}')
            if os.path.isfile(path) and path.endswith('.rpy'):
                if path not in self.merge_listbox.get(0, tk.END):
                    self.merge_listbox.insert(tk.END, path)
            elif os.path.isdir(path):
                rpy_files = glob.glob(os.path.join(path, "**", "*.rpy"), recursive=True)
                for f in rpy_files:
                    if f not in self.merge_listbox.get(0, tk.END):
                        self.merge_listbox.insert(tk.END, f)

    def browse_excel_file(self):
        f = filedialog.askopenfilename(
            title=self.t("dlg_select_excel"),
            filetypes=[("Excel Files", "*.xlsx *.xls")]
        )
        if f:
            self.entry_excel.delete(0, tk.END)
            self.entry_excel.insert(0, f)

    def add_merge_rpy_files(self):
        files = filedialog.askopenfilenames(
            title=self.t("dlg_select_rpy"),
            filetypes=[("Ren'Py Script", "*.rpy"), ("All Files", "*.*")]
        )
        for f in files:
            if f not in self.merge_listbox.get(0, tk.END):
                self.merge_listbox.insert(tk.END, f)

    def add_merge_rpy_folder(self):
        d = filedialog.askdirectory(title=self.t("dlg_select_folder"))
        if d:
            rpy_files = glob.glob(os.path.join(d, "**", "*.rpy"), recursive=True)
            for f in rpy_files:
                if f not in self.merge_listbox.get(0, tk.END):
                    self.merge_listbox.insert(tk.END, f)

    def process_merge(self):
        excel_file = self.entry_excel.get().strip()
        target_rpy_files = list(self.merge_listbox.get(0, tk.END))

        if not excel_file or not os.path.exists(excel_file):
            messagebox.showwarning(self.t("warning"), self.t("msg_warn_excel"))
            return

        if not target_rpy_files:
            messagebox.showwarning(self.t("warning"), self.t("msg_warn_rpy_folder"))
            return

        rpy_file_map = {os.path.basename(f): f for f in target_rpy_files}

        possible_translated_cols = {"Translated", "Bản dịch", "Traducción", "译文", "Übersetzung", "Traduction", "Terjemahan", "الترجمة", "Çeviri", "Tradução", "Tłumaczenie", "Перевод", "Переклад", "翻訳", "번역"}
        possible_original_cols = {"Original", "Bản gốc", "Văn bản gốc", "原文", "Texto Original", "Oryginał", "Оригинал", "Оригінал", "원문"}
        possible_file_cols = {"File", "Tên File", "Archivo", "文件名", "Datei", "Fichier", "Nama Berkas", "اسم الملف", "Dosya Adı", "Nome do Arquivo", "Nazwa Pliku", "Имя файла", "Назва файлу", "ファイル名", "파일명", "Nome File"}
        possible_label_cols = {"Label", "Nhãn (Label)", "Etiqueta", "标签", "Étiquette", "Rótulo", "Etykieta", "Метка", "Мітка", "ラベル", "라벨"}
        possible_notes_sheets = {"Notes", "Ghi chú", "Lưu ý!", "Notas", "备注", "Notizen", "Catatan", "ملاحظات", "Notlar", "Notatki", "Заметки", "Примітки", "メモ", "메모"}
        
        # Danh sách tất cả các tên sheet Bổ sung trong 16 ngôn ngữ
        possible_addition_sheets = {trans.get("sheet_addition_name", "").strip() for trans in TRANSLATIONS.values()}
        possible_addition_sheets.update({"Addition", "Bổ sung", "Adicional", "补充", "Zusatz", "Tambahan", "إضافة", "Ek", "Adição", "Dodatek", "Дополнение", "Доповнення", "追加", "추가", "Aggiunta"})

        try:
            excel_sheets = pd.read_excel(excel_file, sheet_name=None, engine="openpyxl")
            
            dialogue_trans = {}
            string_trans = {}
            addition_trans = {} # Cấu trúc: { file_name: [ (original_text, translated_text), ... ] }

            for sheet_name, df in excel_sheets.items():
                sheet_name_clean = sheet_name.strip()
                if sheet_name_clean in possible_notes_sheets:
                    continue

                df.columns = [str(col).strip() for col in df.columns]

                file_col = next((c for c in df.columns if c in possible_file_cols), df.columns[0])
                label_col = next((c for c in df.columns if c in possible_label_cols), df.columns[3] if len(df.columns) > 3 else None)
                orig_col = next((c for c in df.columns if c in possible_original_cols), df.columns[5] if len(df.columns) > 5 else None)
                trans_col = next((c for c in df.columns if c in possible_translated_cols), df.columns[6] if len(df.columns) > 6 else None)

                is_addition_sheet = sheet_name_clean in possible_addition_sheets

                for _, row in df.iterrows():
                    file_name = str(row.get(file_col, "")).strip()
                    if not file_name or file_name == "nan":
                        if not is_addition_sheet:
                            file_name = f"{sheet_name}.rpy" if not sheet_name.endswith(".rpy") else sheet_name
                        else:
                            continue

                    label_val = str(row.get(label_col, "")).strip() if label_col else ""
                    if label_val == "nan":
                        label_val = ""

                    orig_val = row.get(orig_col, "") if orig_col else ""
                    original_text = "" if pd.isna(orig_val) else str(orig_val).strip()

                    trans_val = row.get(trans_col, "") if trans_col else ""
                    translated_text = "" if pd.isna(trans_val) else str(trans_val).strip()

                    if not original_text or not translated_text:
                        continue

                    if original_text.startswith('"') and original_text.endswith('"') and len(original_text) >= 2:
                        original_text = original_text[1:-1]
                    if translated_text.startswith('"') and translated_text.endswith('"') and len(translated_text) >= 2:
                        translated_text = translated_text[1:-1]

                    if is_addition_sheet:
                        if not file_name.endswith(".rpy"):
                            file_name += ".rpy"
                        if file_name not in addition_trans:
                            addition_trans[file_name] = []
                        addition_trans[file_name].append((original_text, translated_text))
                    else:
                        if label_val == "strings" or not label_val:
                            string_trans[(file_name, original_text)] = translated_text
                        else:
                            dialogue_trans[(file_name, label_val, original_text)] = translated_text

            updated_count = 0
            all_target_files = set(rpy_file_map.keys())

            for file_name in all_target_files:
                file_path = rpy_file_map.get(file_name)
                if not file_path or not os.path.exists(file_path):
                    continue

                with open(file_path, "r", encoding="utf-8") as f:
                    lines = f.readlines()

                new_lines = list(lines)
                is_file_modified = False

                # 1. MERGE THOẠI & STRINGS CÓ SẴN
                current_label = ""
                i = 0
                while i < len(new_lines):
                    line_strip = new_lines[i].strip()

                    if line_strip.startswith("translate ") and line_strip.endswith(":"):
                        parts = line_strip.split()
                        if len(parts) >= 3:
                            current_label = parts[2].replace(":", "").strip()

                    elif line_strip.startswith("#") and '"' in line_strip and current_label:
                        first_q = line_strip.find('"')
                        last_q = line_strip.rfind('"')
                        if first_q != last_q:
                            orig_text = line_strip[first_q + 1 : last_q]
                            trans_text = dialogue_trans.get((file_name, current_label, orig_text))
                            if trans_text:
                                for offset in range(1, 4):
                                    next_idx = i + offset
                                    if next_idx >= len(new_lines):
                                        break
                                    target_line = new_lines[next_idx]
                                    target_strip = target_line.strip()
                                    if target_strip.startswith("#") or target_strip.startswith("translate"):
                                        break
                                    if '"' in target_strip:
                                        t_start = target_line.find('"')
                                        t_end = target_line.rfind('"')
                                        if t_start != t_end:
                                            prefix = target_line[:t_start + 1]
                                            suffix = target_line[t_end:]
                                            new_lines[next_idx] = f"{prefix}{trans_text}{suffix}"
                                            is_file_modified = True
                                            break

                    elif line_strip.startswith("old") and '"' in line_strip:
                        first_q = line_strip.find('"')
                        last_q = line_strip.rfind('"')
                        if first_q != last_q:
                            orig_text = line_strip[first_q + 1 : last_q]
                            trans_text = string_trans.get((file_name, orig_text))
                            if not trans_text:
                                for (fname, otext), ttext in string_trans.items():
                                    if otext == orig_text:
                                        trans_text = ttext
                                        break
                            if trans_text:
                                if i + 1 < len(new_lines) and new_lines[i + 1].strip().startswith("new"):
                                    next_line = new_lines[i + 1]
                                    t_start = next_line.find('"')
                                    t_end = next_line.rfind('"')
                                    if t_start != t_end:
                                        prefix = next_line[:t_start + 1]
                                        suffix = next_line[t_end:]
                                        new_lines[i + 1] = f"{prefix}{trans_text}{suffix}"
                                        is_file_modified = True
                                else:
                                    indent = new_lines[i][:new_lines[i].find("old")]
                                    new_lines.insert(i + 1, f'{indent}new "{trans_text}"\n')
                                    is_file_modified = True
                    i += 1

                # 2. CHÈN BỔ SUNG TỪ SHEET ADDITION VÀO CUỐI FILE (+2 DÒNG) - NHẬN DIỆN NGÔN NGỮ ĐỘNG
                if file_name in addition_trans:
                    adds = addition_trans[file_name]
                    if adds:
                        detected_lang = ""
                        for line in lines:
                            line_s = line.strip()
                            if line_s.startswith("translate ") and ":" in line_s:
                                parts = line_s.split()
                                if len(parts) >= 3 and parts[1] != "strings":
                                    detected_lang = parts[1]
                                    break
                                elif len(parts) >= 2 and parts[1] == "strings":
                                    break

                        while len(new_lines) > 0 and new_lines[-1].strip() == "":
                            new_lines.pop()
                        
                        new_lines.append("\n\n")
                        
                        if detected_lang:
                            new_lines.append(f"translate {detected_lang} strings:\n")
                        else:
                            new_lines.append("translate strings:\n")
                            
                        for orig_text, trans_text in adds:
                            new_lines.append(f'    old "{orig_text}"\n')
                            new_lines.append(f'    new "{trans_text}"\n')
                        
                        is_file_modified = True

                if is_file_modified:
                    with open(file_path, "w", encoding="utf-8") as f:
                        f.writelines(new_lines)
                    updated_count += 1

            messagebox.showinfo(self.t("success"), self.t("msg_success_merge").format(updated_count))

        except Exception as e:
            messagebox.showerror(self.t("error"), self.t("msg_err_merge").format(str(e)))

if __name__ == "__main__":
    if HAS_DND:
        root = TkinterDnD.Tk()
    else:
        root = tk.Tk()
    app = RPYtoEXCELApp(root)
    root.mainloop()